"""
RRWWTF Master Historical Dataset — Step 1 (Restart)

Builds ONE master daily dataset (Date, Total_Treated_MGD, Rainfall_in) from
the RRWWTF Daily Effluent Log Excel workbooks in raw/2.7.7 Wastewater Flows
and Runtimes/{2018,2019,2020,2021,2022,2023_partial,2024}/.

Extraction logic (sheet name, header detection, day-position row mapping,
month/year parsing from sheet title or filename) is the same proven logic
already used and validated in extract_flow_data.py / step2_daily_flow_rainfall.py /
step2_rainfall_integration.py. It has been re-verified here against the raw
workbooks before use (see conversation / manual inspection).

Rainfall is the plant-recorded "Rain" column from the same Daily Effluent Log
sheet used for flow. No external weather source is used.

This script performs extraction, combination, and VALIDATION ONLY. No
statistics (mean/median/std/percentiles), no correlation/lag analysis, no
outlier removal, no imputation, no visualization, and no modeling are
performed here.

Usage:
    python scripts/create_master_dataset.py
"""

import calendar
import re
from dataclasses import dataclass, field
from pathlib import Path

import openpyxl
import pandas as pd

# ---------------------------------------------------------------------------
# Configuration
# ---------------------------------------------------------------------------

BASE_DIR = Path(__file__).resolve().parent.parent
RAW_DIR = BASE_DIR / "raw" / "2.7.7 Wastewater Flows and Runtimes"
OUTPUT_DIR = BASE_DIR / "output"

YEAR_DIRS = ["2018", "2019", "2020", "2021", "2022", "2023_partial", "2024"]

MASTER_DIR = OUTPUT_DIR / "00_master_dataset"
CSV_PATH = MASTER_DIR / "richmond_master_flow_rainfall_dataset.csv"
XLSX_PATH = MASTER_DIR / "richmond_master_flow_rainfall_dataset.xlsx"

SHEET_NAME_KEYWORD = "daily effluent"
DATE_HEADER = "date"
MGD_HEADER = "total treated mgd"
RAIN_HEADER = "rain"

MONTH_ABBR_TO_NUM = {
    "jan": 1, "feb": 2, "mar": 3, "apr": 4, "may": 5, "jun": 6,
    "jul": 7, "aug": 8, "sep": 9, "oct": 10, "nov": 11, "dec": 12,
}


@dataclass
class FileResult:
    path: Path
    status: str = "failed"  # "ok" or "failed"
    reason: str = ""
    n_rows: int = 0
    notes: list = field(default_factory=list)


# ---------------------------------------------------------------------------
# Discovery
# ---------------------------------------------------------------------------

def discover_workbooks(raw_dir: Path) -> list[Path]:
    """Find RRWWTF Excel workbooks only in the specified yearly folders."""
    files = []
    for year_dir_name in YEAR_DIRS:
        year_dir = raw_dir / year_dir_name
        if not year_dir.exists():
            continue
        for p in sorted(year_dir.glob("*.xlsx")):
            if p.name.startswith("~$"):  # Excel lock files
                continue
            files.append(p)
    return files


# ---------------------------------------------------------------------------
# Workbook structure parsing helpers (same proven logic as extract_flow_data.py
# and step2_daily_flow_rainfall.py, re-verified against source workbooks)
# ---------------------------------------------------------------------------

def month_num_from_text(text: str):
    m = re.search(r"[A-Za-z]{3,}", text)
    if not m:
        return None
    return MONTH_ABBR_TO_NUM.get(m.group(0)[:3].lower())


def parse_month_year_from_sheet(ws, max_scan_rows: int = 10):
    """Look for a 'Month of <Month>, <Year>' title near the top of the sheet."""
    for r in range(1, max_scan_rows + 1):
        val = ws.cell(row=r, column=1).value
        if isinstance(val, str) and "month of" in val.lower():
            m = re.search(r"month of\s+([A-Za-z]+)\s*,?\s*(\d{4})", val, re.IGNORECASE)
            if m:
                month_num = month_num_from_text(m.group(1))
                year = int(m.group(2))
                if month_num:
                    return month_num, year
    return None, None


def parse_month_year_from_filename(path: Path):
    """Fallback: infer month/year from the workbook filename."""
    name = path.stem
    year_match = re.search(r"(19|20)\d{2}", name)
    year = int(year_match.group(0)) if year_match else None
    month_num = None
    for token in re.findall(r"[A-Za-z]+", name):
        mn = MONTH_ABBR_TO_NUM.get(token[:3].lower())
        if mn:
            month_num = mn
            break
    return month_num, year


def find_daily_effluent_sheet(wb):
    for sn in wb.sheetnames:
        if SHEET_NAME_KEYWORD in sn.lower():
            return sn
    return None


def find_header_row_and_columns(ws, max_scan_rows: int = 15):
    """Locate the header row containing Date, Total Treated MGD, and Rain."""
    max_col = ws.max_column or 1
    for r in range(1, max_scan_rows + 1):
        row_vals = [ws.cell(row=r, column=c).value for c in range(1, max_col + 1)]
        norm = [str(v).strip().lower() if v is not None else "" for v in row_vals]
        if DATE_HEADER in norm and any(MGD_HEADER in v for v in norm):
            date_col = norm.index(DATE_HEADER) + 1
            mgd_col = next(i for i, v in enumerate(norm) if MGD_HEADER in v) + 1
            rain_col = norm.index(RAIN_HEADER) + 1 if RAIN_HEADER in norm else None
            return r, date_col, mgd_col, rain_col
    return None, None, None, None


def coerce_numeric(raw_value, row_num: int, notes: list, field_label: str):
    """Convert a raw cell value to float, or None if blank/non-numeric."""
    if isinstance(raw_value, bool):
        return None
    if isinstance(raw_value, (int, float)):
        return float(raw_value)
    if isinstance(raw_value, str):
        stripped = raw_value.strip()
        if not stripped:
            return None
        try:
            return float(stripped)
        except ValueError:
            notes.append(f"Row {row_num}: non-numeric {field_label} value {raw_value!r} treated as missing")
            return None
    return None


# ---------------------------------------------------------------------------
# Process a single workbook into daily master records
# ---------------------------------------------------------------------------

def process_workbook(path: Path, source_year_dir: str):
    result = FileResult(path=path)

    try:
        wb = openpyxl.load_workbook(path, data_only=True, read_only=False)
    except Exception as e:
        result.reason = f"Could not open workbook: {e}"
        return result, []

    sheet_name = find_daily_effluent_sheet(wb)
    if sheet_name is None:
        result.reason = "No 'Daily Effluent Log' sheet found in workbook (not a daily log file)"
        wb.close()
        return result, []

    ws = wb[sheet_name]
    header_row, date_col, mgd_col, rain_col = find_header_row_and_columns(ws)
    if header_row is None:
        result.reason = "Could not locate 'Date' / 'Total Treated MGD' header row in Daily Effluent Log sheet"
        wb.close()
        return result, []
    if rain_col is None:
        result.notes.append("No 'Rain' column found in header row; Rainfall_in will be NaN for this file")

    month_num, year = parse_month_year_from_sheet(ws)
    month_year_source = "sheet title"
    if month_num is None or year is None:
        month_num, year = parse_month_year_from_filename(path)
        month_year_source = "filename"

    if month_num is None or year is None:
        result.reason = "Could not determine month/year from sheet title or filename"
        wb.close()
        return result, []

    if month_year_source == "filename":
        result.notes.append("Month/year determined from filename (sheet title missing or unparsable)")

    days_in_month = calendar.monthrange(year, month_num)[1]
    records = []
    for day in range(1, days_in_month + 1):
        r = header_row + day
        date_cell_val = ws.cell(row=r, column=date_col).value
        mgd_val = ws.cell(row=r, column=mgd_col).value
        rain_val = ws.cell(row=r, column=rain_col).value if rain_col is not None else None

        if isinstance(date_cell_val, (int, float)) and not isinstance(date_cell_val, bool):
            if int(date_cell_val) != day:
                result.notes.append(
                    f"Row {r}: 'Date' column value {date_cell_val} does not match expected "
                    f"day-of-month {day} (record kept, position used for calendar date)"
                )

        mgd = coerce_numeric(mgd_val, r, result.notes, "Total Treated MGD")
        rainfall_in = coerce_numeric(rain_val, r, result.notes, "Rain")
        the_date = pd.Timestamp(year=year, month=month_num, day=day)

        records.append({
            "Date": the_date,
            "Total_Treated_MGD": mgd,
            "Rainfall_in": rainfall_in,
            "Year": the_date.year,
            "Month": the_date.month,
            "Source_File": path.name,
            "Source_Year": source_year_dir,
        })

    wb.close()
    result.status = "ok"
    result.n_rows = len(records)
    return result, records


# ---------------------------------------------------------------------------
# Duplicate date handling
# ---------------------------------------------------------------------------

def analyze_duplicates(df: pd.DataFrame):
    """
    Identify duplicate dates. Exact duplicates (same Total_Treated_MGD and
    Rainfall_in, including NaN==NaN treated as equal) are flagged for removal.
    Conflicting duplicates (differing values) are kept in the master dataset
    and reported separately for manual review.
    """
    dup_report_rows = []
    exact_dup_removed_rows = []
    indices_to_drop = []

    date_counts = df["Date"].value_counts()
    dup_dates = sorted(date_counts[date_counts > 1].index)

    for d in dup_dates:
        sub = df[df["Date"] == d]
        # Determine if all rows for this date are identical on flow+rainfall
        key_cols = sub[["Total_Treated_MGD", "Rainfall_in"]].apply(
            lambda col: col.fillna("__NaN__")
        )
        is_exact = key_cols.drop_duplicates().shape[0] == 1

        for idx, row in sub.iterrows():
            dup_report_rows.append({
                "Date": row["Date"],
                "Source_File": row["Source_File"],
                "Total_Treated_MGD": row["Total_Treated_MGD"],
                "Rainfall_in": row["Rainfall_in"],
                "Duplicate_Type": "Exact" if is_exact else "Conflicting",
            })

        if is_exact:
            # Keep the first occurrence, mark the rest for removal
            drop_idx = list(sub.index)[1:]
            indices_to_drop.extend(drop_idx)
            for idx in drop_idx:
                r = df.loc[idx]
                exact_dup_removed_rows.append({
                    "Date": r["Date"],
                    "Source_File": r["Source_File"],
                    "Total_Treated_MGD": r["Total_Treated_MGD"],
                    "Rainfall_in": r["Rainfall_in"],
                })

    dup_report_df = pd.DataFrame(dup_report_rows)
    removed_df = pd.DataFrame(exact_dup_removed_rows)
    return dup_dates, dup_report_df, removed_df, indices_to_drop


# ---------------------------------------------------------------------------
# Missing calendar date handling
# ---------------------------------------------------------------------------

def find_missing_calendar_dates(df: pd.DataFrame) -> pd.DataFrame:
    earliest = df["Date"].min()
    latest = df["Date"].max()
    full_range = pd.date_range(earliest, latest, freq="D")
    present_dates = pd.DatetimeIndex(df["Date"].unique())
    missing = full_range.difference(present_dates)
    return pd.DataFrame({"Missing_Date": missing})


# ---------------------------------------------------------------------------
# Main pipeline
# ---------------------------------------------------------------------------

def main():
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    MASTER_DIR.mkdir(parents=True, exist_ok=True)

    workbooks = discover_workbooks(RAW_DIR)

    all_records = []
    file_results = []
    for path in workbooks:
        source_year_dir = path.parent.name
        result, records = process_workbook(path, source_year_dir)
        file_results.append(result)
        all_records.extend(records)

    n_discovered = len(workbooks)
    n_ok = sum(1 for r in file_results if r.status == "ok")
    n_failed = sum(1 for r in file_results if r.status == "failed")

    df = pd.DataFrame(all_records)
    df = df.sort_values(["Date", "Source_File"]).reset_index(drop=True)

    # --- Duplicate date check ---
    dup_dates, dup_report_df, removed_df, drop_indices = analyze_duplicates(df)

    if dup_dates:
        print(f"DUPLICATE DATES FOUND: {len(dup_dates)}")
        for d in dup_dates:
            sub = df[df["Date"] == d]
            print(f"  Date: {d.date()}")
            for _, row in sub.iterrows():
                print(f"    Source_File={row['Source_File']}  Total_Treated_MGD={row['Total_Treated_MGD']}  "
                      f"Rainfall_in={row['Rainfall_in']}")
        if not removed_df.empty:
            print(f"  Exact duplicate rows removed from master dataset: {len(removed_df)}")
            for _, row in removed_df.iterrows():
                print(f"    Removed: Date={row['Date'].date()} Source_File={row['Source_File']} "
                      f"Total_Treated_MGD={row['Total_Treated_MGD']} Rainfall_in={row['Rainfall_in']}")
        conflicting = dup_report_df[dup_report_df["Duplicate_Type"] == "Conflicting"]
        if not conflicting.empty:
            n_conflicting_dates = conflicting["Date"].nunique()
            print(f"  CONFLICTING duplicate dates (kept in master, needs manual review): {n_conflicting_dates}")
        print()
    else:
        print("DUPLICATE DATES FOUND: 0")
        print()

    if drop_indices:
        df = df.drop(index=drop_indices).reset_index(drop=True)

    # --- Missing calendar date check ---
    missing_dates_df = find_missing_calendar_dates(df)

    # --- Sanity checks (report only, no modification) ---
    negative_flow = df[df["Total_Treated_MGD"] < 0]
    negative_rain = df[df["Rainfall_in"] < 0]

    n_records = len(df)
    earliest = df["Date"].min()
    latest = df["Date"].max()
    n_unique_dates = df["Date"].nunique()
    n_dup_dates = len(dup_dates)
    n_missing_calendar_dates = len(missing_dates_df)
    n_valid_flow = int(df["Total_Treated_MGD"].notna().sum())
    n_missing_flow = int(df["Total_Treated_MGD"].isna().sum())
    n_valid_rain = int(df["Rainfall_in"].notna().sum())
    n_missing_rain = int(df["Rainfall_in"].isna().sum())
    flow_min = df["Total_Treated_MGD"].min()
    flow_max = df["Total_Treated_MGD"].max()
    rain_min = df["Rainfall_in"].min()
    rain_max = df["Rainfall_in"].max()

    # --- Records by year (count only, no statistics) ---
    year_counts = df.groupby("Year").size().sort_index()

    # --- Extraction log ---
    extraction_log_rows = []
    for r in file_results:
        extraction_log_rows.append({
            "Source_File": r.path.name,
            "Status": r.status,
            "Records_Extracted": r.n_rows,
            "Notes": "; ".join(r.notes[:10]) + (f" ... and {len(r.notes) - 10} more" if len(r.notes) > 10 else "")
                     if r.notes else r.reason,
        })
    extraction_log_df = pd.DataFrame(extraction_log_rows)

    # --- Reorder master columns per spec ---
    df = df[["Date", "Total_Treated_MGD", "Rainfall_in", "Year", "Month", "Source_File", "Source_Year"]]

    # --- Save CSV ---
    df.to_csv(CSV_PATH, index=False)

    # --- Save Excel with 4 sheets ---
    with pd.ExcelWriter(XLSX_PATH, engine="openpyxl") as writer:
        df.to_excel(writer, sheet_name="Master_Data", index=False)
        missing_dates_df.to_excel(writer, sheet_name="Missing_Dates", index=False)
        if not dup_report_df.empty:
            dup_report_df.to_excel(writer, sheet_name="Duplicate_Check", index=False)
        else:
            pd.DataFrame({"Note": ["No duplicate dates found."]}).to_excel(
                writer, sheet_name="Duplicate_Check", index=False)
        extraction_log_df.to_excel(writer, sheet_name="Extraction_Log", index=False)

    # -------------------------------------------------------------------
    # Validation summary (console)
    # -------------------------------------------------------------------
    print("=" * 60)
    print("VALIDATION SUMMARY")
    print("=" * 60)
    print(f"1. Workbooks discovered:              {n_discovered}")
    print(f"2. Workbooks successfully processed:  {n_ok}")
    print(f"3. Workbooks that failed:             {n_failed}")
    print("4. Names of failed workbooks:")
    for r in file_results:
        if r.status == "failed":
            print(f"     - {r.path.name}  ({r.reason})")
    print(f"5. Extracted daily records:           {n_records}")
    print(f"6. Earliest date:                     {earliest.date()}")
    print(f"7. Latest date:                       {latest.date()}")
    print(f"8. Unique dates:                      {n_unique_dates}")
    print(f"9. Duplicate dates:                   {n_dup_dates}")
    print(f"10. Missing calendar dates:           {n_missing_calendar_dates}")
    print(f"11. Valid Total_Treated_MGD values:   {n_valid_flow}")
    print(f"12. Missing Total_Treated_MGD values: {n_missing_flow}")
    print(f"13. Valid Rainfall_in values:         {n_valid_rain}")
    print(f"14. Missing Rainfall_in values:       {n_missing_rain}")
    print()
    print("Records by year (validation count only):")
    for year, count in year_counts.items():
        print(f"  {year}: {count} records")
    print()

    print("=" * 60)
    print("SANITY CHECKS")
    print("=" * 60)
    print(f"Negative flow values:     {len(negative_flow)}")
    if not negative_flow.empty:
        for _, row in negative_flow.iterrows():
            print(f"    {row['Date'].date()}  Source_File={row['Source_File']}  Total_Treated_MGD={row['Total_Treated_MGD']}")
    print(f"Negative rainfall values: {len(negative_rain)}")
    if not negative_rain.empty:
        for _, row in negative_rain.iterrows():
            print(f"    {row['Date'].date()}  Source_File={row['Source_File']}  Rainfall_in={row['Rainfall_in']}")
    print(f"Minimum extracted flow (MGD):     {flow_min}")
    print(f"Maximum extracted flow (MGD):     {flow_max}")
    print(f"Minimum extracted rainfall (in):  {rain_min}")
    print(f"Maximum extracted rainfall (in):  {rain_max}")
    print()

    # -------------------------------------------------------------------
    # Final dataframe views
    # -------------------------------------------------------------------
    pd.set_option("display.width", 140)
    pd.set_option("display.max_columns", 20)

    print("=" * 60)
    print("FINAL DATAFRAME COLUMNS")
    print("=" * 60)
    print(list(df.columns))
    print()

    print("=" * 60)
    print("FIRST 10 ROWS")
    print("=" * 60)
    print(df.head(10).to_string(index=False))
    print()

    print("=" * 60)
    print("LAST 10 ROWS")
    print("=" * 60)
    print(df.tail(10).to_string(index=False))
    print()

    # -------------------------------------------------------------------
    # Final concise summary
    # -------------------------------------------------------------------
    print("MASTER DATASET CREATED")
    print()
    print("Files processed:")
    print(f"{n_ok} / {n_discovered}")
    print()
    print("Date range:")
    print(f"{earliest.date()} to {latest.date()}")
    print()
    print("Extracted records:")
    print(f"{n_records}")
    print()
    print("Unique dates:")
    print(f"{n_unique_dates}")
    print()
    print("Missing calendar dates:")
    print(f"{n_missing_calendar_dates}")
    print()
    print("Duplicate dates:")
    print(f"{n_dup_dates}")
    print()
    print("Valid flow observations:")
    print(f"{n_valid_flow}")
    print()
    print("Missing flow observations:")
    print(f"{n_missing_flow}")
    print()
    print("Valid rainfall observations:")
    print(f"{n_valid_rain}")
    print()
    print("Missing rainfall observations:")
    print(f"{n_missing_rain}")
    print()
    print("Flow range:")
    print(f"{flow_min:.3f} to {flow_max:.3f} MGD")
    print()
    print("Rainfall range:")
    print(f"{rain_min:.2f} to {rain_max:.2f} inches")
    print()
    print("CSV saved:")
    print(f"{CSV_PATH.relative_to(BASE_DIR).as_posix()}")
    print()
    print("Excel saved:")
    print(f"{XLSX_PATH.relative_to(BASE_DIR).as_posix()}")


if __name__ == "__main__":
    main()
