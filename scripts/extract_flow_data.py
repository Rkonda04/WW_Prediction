"""
RRWWTF Historical Total Treated MGD Extraction — Step 1

Recursively reads RRWWTF Daily Effluent Log Excel workbooks and extracts the
Date and Total Treated MGD observations exactly as recorded in the source
files. This script performs NO outlier removal, imputation, smoothing,
transformation, or modeling. It only extracts, describes, and visualizes the
historical data as-is.

Usage:
    python scripts/extract_flow_data.py
"""

import calendar
import re
from dataclasses import dataclass, field
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import openpyxl
import pandas as pd

# ---------------------------------------------------------------------------
# Configuration
# ---------------------------------------------------------------------------

BASE_DIR = Path(__file__).resolve().parent.parent
RAW_DIR = BASE_DIR / "raw" / "2.7.7 Wastewater Flows and Runtimes"
OUTPUT_DIR = BASE_DIR / "output"
FIG_DIR = OUTPUT_DIR / "figures"

CSV_PATH = OUTPUT_DIR / "richmond_total_treated_mgd.csv"
REPORT_PATH = OUTPUT_DIR / "Richmond_WWTF_Flow_Data_Statistics_Report.md"

TS_FIG_PATH = FIG_DIR / "historical_daily_flow.png"
HIST_FIG_PATH = FIG_DIR / "flow_distribution_histogram.png"
BOX_FIG_PATH = FIG_DIR / "yearly_distribution_boxplot.png"

# Folders under RAW_DIR to skip entirely (not RRWWTF log data)
EXCLUDED_DIR_NAMES = {"pump curves"}

SHEET_NAME_KEYWORD = "daily effluent"
DATE_HEADER = "date"
MGD_HEADER = "total treated mgd"

MONTH_ABBR_TO_NUM = {
    "jan": 1, "feb": 2, "mar": 3, "apr": 4, "may": 5, "jun": 6,
    "jul": 7, "aug": 8, "sep": 9, "oct": 10, "nov": 11, "dec": 12,
}


@dataclass
class FileResult:
    path: Path
    status: str = "failed"  # "ok" or "failed"
    reason: str = ""
    n_rows: int = 0
    notes: list = field(default_factory=list)


# ---------------------------------------------------------------------------
# Step 1: Discovery
# ---------------------------------------------------------------------------

def discover_workbooks(raw_dir: Path) -> list[Path]:
    """Recursively find RRWWTF Excel workbooks, skipping excluded folders."""
    files = []
    for p in sorted(raw_dir.rglob("*.xlsx")):
        if p.name.startswith("~$"):  # Excel lock files
            continue
        rel_parts = [part.lower() for part in p.relative_to(raw_dir).parts[:-1]]
        if any(part in EXCLUDED_DIR_NAMES for part in rel_parts):
            continue
        files.append(p)
    return files


# ---------------------------------------------------------------------------
# Step 1: Workbook structure inspection helpers
# ---------------------------------------------------------------------------

def month_num_from_text(text: str):
    """Map a month name/abbreviation (with common typos) to its number."""
    m = re.search(r"[A-Za-z]{3,}", text)
    if not m:
        return None
    return MONTH_ABBR_TO_NUM.get(m.group(0)[:3].lower())


def parse_month_year_from_sheet(ws, max_scan_rows: int = 10):
    """Look for a 'Month of <Month>, <Year>' title near the top of the sheet."""
    for r in range(1, max_scan_rows + 1):
        val = ws.cell(row=r, column=1).value
        if isinstance(val, str) and "month of" in val.lower():
            m = re.search(r"month of\s+([A-Za-z]+)\s*,?\s*(\d{4})", val, re.IGNORECASE)
            if m:
                month_num = month_num_from_text(m.group(1))
                year = int(m.group(2))
                if month_num:
                    return month_num, year
    return None, None


def parse_month_year_from_filename(path: Path):
    """Fallback: infer month/year from the workbook filename."""
    name = path.stem
    year_match = re.search(r"(19|20)\d{2}", name)
    year = int(year_match.group(0)) if year_match else None
    month_num = None
    for token in re.findall(r"[A-Za-z]+", name):
        mn = MONTH_ABBR_TO_NUM.get(token[:3].lower())
        if mn:
            month_num = mn
            break
    return month_num, year


def find_daily_effluent_sheet(wb):
    for sn in wb.sheetnames:
        if SHEET_NAME_KEYWORD in sn.lower():
            return sn
    return None


def find_header_row_and_columns(ws, max_scan_rows: int = 15):
    """Locate the header row containing 'Date' and 'Total Treated MGD'."""
    max_col = ws.max_column or 1
    for r in range(1, max_scan_rows + 1):
        row_vals = [ws.cell(row=r, column=c).value for c in range(1, max_col + 1)]
        norm = [str(v).strip().lower() if v is not None else "" for v in row_vals]
        if DATE_HEADER in norm and any(MGD_HEADER in v for v in norm):
            date_col = norm.index(DATE_HEADER) + 1
            mgd_col = next(i for i, v in enumerate(norm) if MGD_HEADER in v) + 1
            return r, date_col, mgd_col
    return None, None, None


def coerce_mgd_value(raw_value, row_num: int, notes: list):
    """Convert a raw cell value to a float, or None if blank/non-numeric."""
    if isinstance(raw_value, bool):
        return None
    if isinstance(raw_value, (int, float)):
        return float(raw_value)
    if isinstance(raw_value, str):
        stripped = raw_value.strip()
        if not stripped:
            return None
        try:
            return float(stripped)
        except ValueError:
            notes.append(f"Row {row_num}: non-numeric Total Treated MGD value {raw_value!r} treated as missing")
            return None
    return None


# ---------------------------------------------------------------------------
# Step 1 + 2: Process a single workbook into daily records
# ---------------------------------------------------------------------------

def process_workbook(path: Path):
    result = FileResult(path=path)
    rel_name = str(path.relative_to(RAW_DIR))

    try:
        wb = openpyxl.load_workbook(path, data_only=True, read_only=False)
    except Exception as e:
        result.reason = f"Could not open workbook: {e}"
        return result, []

    sheet_name = find_daily_effluent_sheet(wb)
    if sheet_name is None:
        result.reason = "No 'Daily Effluent Log' sheet found in workbook"
        wb.close()
        return result, []

    ws = wb[sheet_name]
    header_row, date_col, mgd_col = find_header_row_and_columns(ws)
    if header_row is None:
        result.reason = "Could not locate 'Date' / 'Total Treated MGD' header row in Daily Effluent Log sheet"
        wb.close()
        return result, []

    month_num, year = parse_month_year_from_sheet(ws)
    month_year_source = "sheet title"
    if month_num is None or year is None:
        month_num, year = parse_month_year_from_filename(path)
        month_year_source = "filename"

    if month_num is None or year is None:
        result.reason = "Could not determine month/year from sheet title or filename"
        wb.close()
        return result, []

    if month_year_source == "filename":
        result.notes.append("Month/year determined from filename (sheet title missing or unparsable)")

    days_in_month = calendar.monthrange(year, month_num)[1]
    records = []
    for day in range(1, days_in_month + 1):
        r = header_row + day
        date_cell_val = ws.cell(row=r, column=date_col).value
        mgd_val = ws.cell(row=r, column=mgd_col).value

        if isinstance(date_cell_val, (int, float)) and not isinstance(date_cell_val, bool):
            if int(date_cell_val) != day:
                result.notes.append(
                    f"Row {r}: 'Date' column value {date_cell_val} does not match expected "
                    f"day-of-month {day} (record kept, position used for calendar date)"
                )

        mgd = coerce_mgd_value(mgd_val, r, result.notes)
        the_date = pd.Timestamp(year=year, month=month_num, day=day)
        records.append({
            "Date": the_date,
            "Total_Treated_MGD": mgd,
            "Source_File": rel_name,
        })

    wb.close()
    result.status = "ok"
    result.n_rows = len(records)
    return result, records


# ---------------------------------------------------------------------------
# Step 4/5: Descriptive statistics
# ---------------------------------------------------------------------------

def compute_stats(series: pd.Series) -> dict:
    valid = series.dropna()
    if valid.empty:
        return {k: None for k in
                ["Count", "Min", "P5", "P25", "Median", "Mean", "P75", "P95", "Max", "StdDev", "IQR"]}
    p25 = valid.quantile(0.25)
    p75 = valid.quantile(0.75)
    return {
        "Count": int(valid.count()),
        "Min": valid.min(),
        "P5": valid.quantile(0.05),
        "P25": p25,
        "Median": valid.quantile(0.50),
        "Mean": valid.mean(),
        "P75": p75,
        "P95": valid.quantile(0.95),
        "Max": valid.max(),
        "StdDev": valid.std(),
        "IQR": p75 - p25,
    }


def fmt(v, nd=3):
    if v is None or pd.isna(v):
        return "N/A"
    if isinstance(v, (int,)):
        return str(v)
    return f"{v:,.{nd}f}"


# ---------------------------------------------------------------------------
# Step 6: Visualizations
# ---------------------------------------------------------------------------

def plot_time_series(df: pd.DataFrame, out_path: Path):
    valid = df.dropna(subset=["Total_Treated_MGD"]).sort_values("Date")
    fig, ax = plt.subplots(figsize=(16, 6))
    ax.plot(valid["Date"], valid["Total_Treated_MGD"], marker="o", markersize=2,
            linewidth=0.6, color="#2b6cb0", alpha=0.8)
    ax.set_title("RRWWTF Historical Daily Total Treated Flow (Raw Observations)")
    ax.set_xlabel("Date")
    ax.set_ylabel("Total Treated MGD")
    ax.grid(True, linewidth=0.4, alpha=0.5)
    fig.autofmt_xdate()
    fig.tight_layout()
    fig.savefig(out_path, dpi=150)
    plt.close(fig)


def plot_histogram(df: pd.DataFrame, out_path: Path):
    valid = df["Total_Treated_MGD"].dropna()
    fig, ax = plt.subplots(figsize=(10, 6))
    ax.hist(valid, bins=50, color="#2b6cb0", edgecolor="white", alpha=0.85)
    ax.set_title("Distribution of Daily Total Treated Flow — All Historical Observations")
    ax.set_xlabel("Total Treated MGD")
    ax.set_ylabel("Frequency (Number of Days)")
    ax.grid(True, axis="y", linewidth=0.4, alpha=0.5)
    fig.tight_layout()
    fig.savefig(out_path, dpi=150)
    plt.close(fig)


def plot_yearly_boxplot(df: pd.DataFrame, out_path: Path):
    d = df.dropna(subset=["Total_Treated_MGD"]).copy()
    d["Year"] = d["Date"].dt.year
    years = sorted(d["Year"].unique())
    data = [d.loc[d["Year"] == y, "Total_Treated_MGD"].values for y in years]
    fig, ax = plt.subplots(figsize=(12, 6))
    ax.boxplot(data, tick_labels=[str(y) for y in years], showmeans=False)
    ax.set_title("Yearly Distribution of Daily Total Treated Flow (Raw Observations)")
    ax.set_xlabel("Year")
    ax.set_ylabel("Total Treated MGD")
    ax.grid(True, axis="y", linewidth=0.4, alpha=0.5)
    fig.tight_layout()
    fig.savefig(out_path, dpi=150)
    plt.close(fig)


# ---------------------------------------------------------------------------
# Coverage helpers
# ---------------------------------------------------------------------------

def collapse_to_ranges(dates: pd.DatetimeIndex) -> list[tuple]:
    """Collapse a sorted list of dates into contiguous [start, end] ranges."""
    if len(dates) == 0:
        return []
    dates = sorted(dates)
    ranges = []
    start = prev = dates[0]
    for d in dates[1:]:
        if (d - prev).days == 1:
            prev = d
            continue
        ranges.append((start, prev))
        start = prev = d
    ranges.append((start, prev))
    return ranges


# ---------------------------------------------------------------------------
# Report generation
# ---------------------------------------------------------------------------

def build_report(df, file_results, overall_stats, yearly_stats_df,
                  missing_ranges, dup_table, coverage):
    lines = []
    lines.append("# Richmond Regional WWTF — Historical Total Treated MGD: Data Statistics Report")
    lines.append("")
    lines.append(
        "This report describes the historical **Total Treated MGD** data extracted from the "
        "RRWWTF Daily Effluent Log workbooks exactly as recorded. No outliers were removed, no "
        "values were imputed, and no smoothing or transformation was applied. This is Step 1 of "
        "the wastewater flow analysis project: understanding what data exists before any "
        "analytical or modeling decisions are made."
    )
    lines.append("")

    # 1. Dataset Overview
    lines.append("## 1. Dataset Overview")
    lines.append("")
    lines.append(f"- Workbooks discovered: **{coverage['n_discovered']}**")
    lines.append(f"- Workbooks successfully processed: **{coverage['n_ok']}**")
    lines.append(f"- Workbooks that failed to process: **{coverage['n_failed']}**")
    lines.append(f"- Date range: **{coverage['earliest'].date()} to {coverage['latest'].date()}**")
    lines.append(f"- Total extracted observations (calendar-day rows): **{coverage['n_obs']}**")
    lines.append(f"- Valid Total Treated MGD observations: **{coverage['n_valid']}**")
    lines.append(f"- Missing/null Total Treated MGD observations: **{coverage['n_missing']}**")
    lines.append(f"- Unique dates: **{coverage['n_unique_dates']}**")
    lines.append(f"- Duplicate dates: **{coverage['n_dup_dates']}**")
    lines.append(f"- Missing calendar dates (gaps in coverage, earliest–latest): **{coverage['n_missing_calendar_dates']}**")
    lines.append("")

    # 2. Overall Flow Statistics
    lines.append("## 2. Overall Flow Statistics — Total Treated MGD")
    lines.append("")
    lines.append("| Statistic | Value (MGD) |")
    lines.append("|---|---:|")
    for label, key in [
        ("Count", "Count"), ("Minimum", "Min"), ("5th percentile", "P5"),
        ("25th percentile", "P25"), ("Median (50th percentile)", "Median"),
        ("Mean", "Mean"), ("75th percentile", "P75"), ("95th percentile", "P95"),
        ("Maximum", "Max"), ("Standard deviation", "StdDev"), ("IQR", "IQR"),
    ]:
        v = overall_stats[key]
        lines.append(f"| {label} | {fmt(v) if key != 'Count' else v} |")
    lines.append("")
    lines.append(
        "*Percentiles above describe the shape of the historical distribution only. "
        "No observations were removed or flagged as outliers at this stage.*"
    )
    lines.append("")

    # 3. Statistics by Year
    lines.append("## 3. Statistics by Year — Total Treated MGD")
    lines.append("")
    cols = ["Year", "Count", "Min", "P5", "P25", "Median", "Mean", "P75", "P95", "Max", "StdDev", "IQR", "Missing"]
    lines.append("| " + " | ".join(cols) + " |")
    lines.append("|" + "---|" * len(cols))
    for _, row in yearly_stats_df.iterrows():
        vals = [str(int(row["Year"]))]
        vals += [str(int(row["Count"])) if pd.notna(row["Count"]) else "0"]
        for k in ["Min", "P5", "P25", "Median", "Mean", "P75", "P95", "Max", "StdDev", "IQR"]:
            vals.append(fmt(row[k]))
        vals.append(str(int(row["Missing"])))
        lines.append("| " + " | ".join(vals) + " |")
    lines.append("")

    # 4. Figures
    lines.append("## 4. Figures")
    lines.append("")
    lines.append("### Historical Daily Flow")
    lines.append("Date vs Total Treated MGD, all valid raw observations, no smoothing:")
    lines.append("")
    lines.append(f"![Historical Daily Flow]({TS_FIG_PATH.relative_to(OUTPUT_DIR).as_posix()})")
    lines.append("")
    lines.append("### Flow Distribution")
    lines.append("Histogram of all valid Total Treated MGD observations:")
    lines.append("")
    lines.append(f"![Flow Distribution Histogram]({HIST_FIG_PATH.relative_to(OUTPUT_DIR).as_posix()})")
    lines.append("")
    lines.append("### Yearly Distribution")
    lines.append("Boxplot of Total Treated MGD grouped by year:")
    lines.append("")
    lines.append(f"![Yearly Distribution Boxplot]({BOX_FIG_PATH.relative_to(OUTPUT_DIR).as_posix()})")
    lines.append("")

    # 5. Data Quality Notes
    lines.append("## 5. Data Quality Notes")
    lines.append("")
    lines.append(
        "No observations were removed, corrected, or imputed. The items below are reported "
        "for future investigation only."
    )
    lines.append("")

    lines.append(f"### Missing Total Treated MGD observations: {coverage['n_missing']}")
    lines.append("")
    lines.append(
        "These are calendar days for which a workbook row exists (the date was within a "
        "processed month) but the Total Treated MGD cell was blank or non-numeric in the "
        "source file."
    )
    lines.append("")

    lines.append(f"### Missing calendar dates: {coverage['n_missing_calendar_dates']}")
    lines.append("")
    lines.append("Calendar days between the earliest and latest observation with **no workbook row at all** "
                  "(i.e., no source file covered that month/date):")
    lines.append("")
    if missing_ranges:
        lines.append("| Start | End | Days |")
        lines.append("|---|---|---:|")
        for start, end in missing_ranges:
            n_days = (end - start).days + 1
            lines.append(f"| {start.date()} | {end.date()} | {n_days} |")
    else:
        lines.append("None — every calendar date in range has at least one workbook row.")
    lines.append("")

    lines.append(f"### Duplicate dates: {coverage['n_dup_dates']}")
    lines.append("")
    if not dup_table.empty:
        lines.append("| Date | Count | Source Files |")
        lines.append("|---|---:|---|")
        for _, row in dup_table.iterrows():
            lines.append(f"| {row['Date'].date()} | {row['Count']} | {row['Source_Files']} |")
    else:
        lines.append("None — no date appears more than once in the extracted dataset.")
    lines.append("")

    lines.append("### Workbooks that could not be processed")
    lines.append("")
    failed = [r for r in file_results if r.status == "failed"]
    if failed:
        lines.append("| File | Reason |")
        lines.append("|---|---|")
        for r in failed:
            lines.append(f"| {r.path.relative_to(RAW_DIR).as_posix()} | {r.reason} |")
    else:
        lines.append("None — all discovered workbooks were processed.")
    lines.append("")

    lines.append("### Structural inconsistencies encountered while reading files")
    lines.append("")
    any_notes = False
    for r in file_results:
        if r.status == "ok" and r.notes:
            any_notes = True
            lines.append(f"**{r.path.relative_to(RAW_DIR).as_posix()}**")
            for n in r.notes[:20]:
                lines.append(f"- {n}")
            if len(r.notes) > 20:
                lines.append(f"- ... and {len(r.notes) - 20} more note(s) for this file")
            lines.append("")
    if not any_notes:
        lines.append("None recorded.")
        lines.append("")

    return "\n".join(lines)


# ---------------------------------------------------------------------------
# Main pipeline
# ---------------------------------------------------------------------------

def main():
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    FIG_DIR.mkdir(parents=True, exist_ok=True)

    workbooks = discover_workbooks(RAW_DIR)

    all_records = []
    file_results = []
    for path in workbooks:
        result, records = process_workbook(path)
        file_results.append(result)
        all_records.extend(records)

    n_ok = sum(1 for r in file_results if r.status == "ok")
    n_failed = sum(1 for r in file_results if r.status == "failed")

    df = pd.DataFrame(all_records)
    df = df.sort_values(["Date", "Source_File"]).reset_index(drop=True)
    df.to_csv(CSV_PATH, index=False)

    # --- Coverage statistics ---
    earliest = df["Date"].min()
    latest = df["Date"].max()
    n_obs = len(df)
    n_valid = int(df["Total_Treated_MGD"].notna().sum())
    n_missing = int(df["Total_Treated_MGD"].isna().sum())

    date_counts = df["Date"].value_counts()
    n_unique_dates = int(date_counts.shape[0])
    dup_dates = date_counts[date_counts > 1].sort_index()
    n_dup_dates = int(dup_dates.shape[0])

    full_range = pd.date_range(earliest, latest, freq="D")
    present_dates = pd.DatetimeIndex(df["Date"].unique())
    missing_calendar_dates = full_range.difference(present_dates)
    n_missing_calendar_dates = int(len(missing_calendar_dates))
    missing_ranges = collapse_to_ranges(missing_calendar_dates)

    dup_rows = []
    for d in dup_dates.index:
        files = sorted(df.loc[df["Date"] == d, "Source_File"].unique())
        dup_rows.append({"Date": d, "Count": int(dup_dates.loc[d]), "Source_Files": "; ".join(files)})
    dup_table = pd.DataFrame(dup_rows)

    coverage = {
        "n_discovered": len(workbooks),
        "n_ok": n_ok,
        "n_failed": n_failed,
        "earliest": earliest,
        "latest": latest,
        "n_obs": n_obs,
        "n_valid": n_valid,
        "n_missing": n_missing,
        "n_unique_dates": n_unique_dates,
        "n_dup_dates": n_dup_dates,
        "n_missing_calendar_dates": n_missing_calendar_dates,
    }

    # --- Descriptive statistics ---
    overall_stats = compute_stats(df["Total_Treated_MGD"])

    yearly_rows = []
    d2 = df.copy()
    d2["Year"] = d2["Date"].dt.year
    for year, g in d2.groupby("Year"):
        s = compute_stats(g["Total_Treated_MGD"])
        s["Year"] = year
        s["Missing"] = int(g["Total_Treated_MGD"].isna().sum())
        yearly_rows.append(s)
    yearly_stats_df = pd.DataFrame(yearly_rows).sort_values("Year").reset_index(drop=True)
    yearly_stats_df.to_csv(OUTPUT_DIR / "richmond_total_treated_mgd_yearly_stats.csv", index=False)

    # --- Plots ---
    plot_time_series(df, TS_FIG_PATH)
    plot_histogram(df, HIST_FIG_PATH)
    plot_yearly_boxplot(df, BOX_FIG_PATH)

    # --- Report ---
    report_text = build_report(df, file_results, overall_stats, yearly_stats_df,
                                missing_ranges, dup_table, coverage)
    REPORT_PATH.write_text(report_text, encoding="utf-8")

    # --- Console summary ---
    print("RRWWTF FLOW DATA SUMMARY")
    print(f"Workbooks discovered:            {coverage['n_discovered']}")
    print(f"Workbooks processed successfully: {coverage['n_ok']}")
    print(f"Workbooks failed:                {coverage['n_failed']}")
    print(f"Date range:                      {earliest.date()} to {latest.date()}")
    print(f"Total observations:              {n_obs}")
    print(f"Valid flow observations:         {n_valid}")
    print(f"Missing flow observations:       {n_missing}")
    print(f"Unique dates:                    {n_unique_dates}")
    print(f"Duplicate dates:                 {n_dup_dates}")
    print(f"Missing calendar dates:          {n_missing_calendar_dates}")
    print(f"Mean flow (MGD):                 {fmt(overall_stats['Mean'])}")
    print(f"Median flow (MGD):                {fmt(overall_stats['Median'])}")
    print(f"5th percentile (MGD):             {fmt(overall_stats['P5'])}")
    print(f"95th percentile (MGD):            {fmt(overall_stats['P95'])}")
    print(f"Minimum (MGD):                    {fmt(overall_stats['Min'])}")
    print(f"Maximum (MGD):                    {fmt(overall_stats['Max'])}")
    print()
    print("Output files:")
    print(f"  {CSV_PATH}")
    print(f"  {OUTPUT_DIR / 'richmond_total_treated_mgd_yearly_stats.csv'}")
    print(f"  {REPORT_PATH}")
    print(f"  {TS_FIG_PATH}")
    print(f"  {HIST_FIG_PATH}")
    print(f"  {BOX_FIG_PATH}")


if __name__ == "__main__":
    main()
